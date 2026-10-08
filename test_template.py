# -*- coding: utf-8 -*-
"""模板自检：写入边界用例 -> 用 formulas 引擎真算一遍 -> 断言结果。

用法: py test_template.py
"""
import os
import sys

import formulas
from openpyxl import load_workbook

import build_template

TMP = "_selftest.xlsx"

# ---------- 用例 ----------
# 初选表: (粉丝 G, 加购 F) -> 期望档位
初选用例 = [
    (300, 0, "低粉爆款"), (300, 9999, "低粉爆款"),
    (5000, 900, "大爆款"), (499, 900, "低粉爆款"),
    (1000, 500, "常规爆款"), (1999, 401, "常规爆款"),
    (2000, 500, "未达标"), (5000, 100, "未达标"),
    (1000, 400, "未达标"), (1000, 800, "常规爆款"),
    # 文本不能当成「很大」的数字（Excel 里文本>数字恒为真）
    (1000, "?", ""), ("?", 500, ""), (None, 500, ""), (1000, None, ""),
]
# 二选表: (笔记总数 C, 评价数 D, 拍摄 F, 超越 G) -> (权重 E, 结论 H)
二选用例 = [
    (1500, 1200, "是", "是", "高权重", "进调研"),
    (1500, 500, "是", "是", "低权重", "进调研"),
    (1000, 1000, "是", "是", "低权重", "进调研"),
    (2000, 3000, "是", "否", "高权重", "放弃"),
    (2000, 3000, "否", "是", "高权重", "放弃"),
    (None, None, None, None, "", ""),
    # 新增：填一半时的中间态，不能误报「放弃」
    (343, None, "是", None, "低权重", ""),        # 只填了 F，G 待填
    (343, None, "否", None, "低权重", "放弃"),       # F=否 立即放弃
    (343, None, None, "是", "低权重", ""),        # 只填了 G，F 待填
    (343, None, "否", "否", "低权重", "放弃"),
    (343, None, "否", "是", "低权重", "放弃"),       # F=否 优先于 G=是
    # 笔记数未知时填「?」：必须留空，不能变成高权重
    ("?", None, None, None, "", ""),
    ("?", "?", None, None, "", ""),
    (155, "?", None, None, "低权重", ""),             # 155<=1000 已能定论
    (2000, "?", None, None, "", ""),                 # 2000>1000 但评价数未知
    (2000, 1500, "是", "是", "高权重", "进调研"),
]
# 调研表: (售价 C, 批发 D, 运费 E, 包装 F, 加购 I, 高粉 M, 低粉 N, 爆文天数 P)
调研用例 = [
    # 单件利润 每日利润 梯队 竞品总数 入场判断
    (100, 30, 5, 2, 1000, 3, 1, 40, 58, 58000, "第一梯队", 4, "可以入场"),
    (100, 30, 5, 2, 1000, 5, 2, 40, 58, 58000, "第一梯队", 7, "放弃·生态位占满"),
    (100, 30, 5, 2, 100, 2, 0, 45, 58, 5800, "第二梯队", 2, "谨慎入场"),
    (100, 30, 5, 2, 20, 1, 1, 7, 58, 1160, "第三梯队", 2, "可以入场"),
    (100, 30, 5, 2, 20, 1, 1, 20, 58, 1160, "第三梯队", 2, "待判断·爆文15~30天"),
    (100, 30, 5, 2, 20, 1, 0, None, 58, 1160, "第三梯队", 1, "待查首篇爆文时间"),
    (100, 30, 5, 2, 10, 0, 0, 5, 58, 580, "不达标", 0, "可以入场"),
    # 文本输入不能崩、也不能被当成大数
    (100, 30, 5, 2, "?", 3, 1, 40, 58, "", "", 4, "可以入场"),
    (100, 30, 5, 2, 1000, "?", "?", 40, 58, 58000, "第一梯队", "", ""),
    ("?", 30, 5, 2, 1000, 3, 1, 40, "", "", "", 4, "可以入场"),
    # 批发价没填（非数字）时不能算利潦 —— 否则等于「货不要钱」
    (100, "?", 5, 2, 1000, 3, 1, 40, "", "", "", 4, "可以入场"),
    (100, None, 5, 2, 1000, 3, 1, 40, "", "", "", 4, "可以入场"),
    (28.8, 8, None, None, 400, 0, 3, 100, 19.36, 7744, "第二梯队", 3, "可以入场"),
]


def main():
    build_template.build(TMP)

    wb = load_workbook(TMP)
    cv, er, dy = wb["初选表"], wb["二选表"], wb["调研表"]

    for i, (fans, cart, _) in enumerate(初选用例, start=2):
        cv.cell(row=i, column=6, value=cart)
        cv.cell(row=i, column=7, value=fans)
    for i, (notes, cmts, shoot, beat, _, _) in enumerate(二选用例, start=2):
        er.cell(row=i, column=3, value=notes)
        er.cell(row=i, column=4, value=cmts)
        er.cell(row=i, column=6, value=shoot)
        er.cell(row=i, column=7, value=beat)
    for i, row in enumerate([r[:8] for r in 调研用例], start=2):
        price, whole, ship, pack, cart, hi, lo, days = row
        dy.cell(row=i, column=3, value=price)
        dy.cell(row=i, column=4, value=whole)
        dy.cell(row=i, column=5, value=ship)
        dy.cell(row=i, column=6, value=pack)
        dy.cell(row=i, column=9, value=cart)
        dy.cell(row=i, column=13, value=hi)
        dy.cell(row=i, column=14, value=lo)
        dy.cell(row=i, column=16, value=days)
    wb.save(TMP)

    print("用 formulas 引擎重算中（首次较慢）…")
    sol = formulas.ExcelModel().loads(TMP).finish().calculate()

    def val(sheet, cell):
        suffix = f"]{sheet}'!{cell}".upper()
        for k, v in sol.items():
            if k.upper().endswith(suffix):
                try:
                    return v.value[0, 0]
                except Exception:
                    return v
        return "<<未找到>>"

    fails = []

    def check(label, got, want):
        if isinstance(got, (int, float)) and isinstance(want, (int, float)):
            ok = abs(float(got) - float(want)) < 1e-6
        else:
            ok = str(got) == str(want)
        if not ok:
            fails.append(f"{label}: 得到 {got!r}，期望 {want!r}")
        print(f"  {'OK ' if ok else 'FAIL'} {label:<34} {got!r}")

    print("\n初选表 档位")
    for i, (fans, cart, want) in enumerate(初选用例, start=2):
        check(f"粉丝{fans} 加购{cart}", val("初选表", f"H{i}"), want)

    print("\n二选表 权重 / 结论")
    for i, (_n, _c, _s, _b, w_want, j_want) in enumerate(二选用例, start=2):
        check(f"row{i} 权重", val("二选表", f"E{i}"), w_want)
        check(f"row{i} 结论", val("二选表", f"H{i}"), j_want)

    print("\n调研表 利润 / 梯队 / 竞品总数 / 入场判断")
    for i, case in enumerate(调研用例, start=2):
        want = case[8:]
        for col, w in zip(("H", "J", "K", "O", "Q"), want):
            check(f"row{i} {col}", val("调研表", f"{col}{i}"), w)

    print()
    if fails:
        print(f"❌ {len(fails)} 项不符：")
        for f in fails:
            print("   -", f)
        sys.exit(1)
    print("✅ 全部通过")


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    main()
