# -*- coding: utf-8 -*-
"""
银河选品广场（aiyinhe.art）抓取 + 按《小红书选品 SOP 4.0》分档

数据源: GET https://aiyinhe.art/api/plaza  (需要登录态 mc_token)
        未登录只有 /api/public/plaza，20 条脱敏预览（无粉丝数、加购取整），不能用。

SOP 分档规则（与 filter_scrape.py 完全一致）:
  低粉爆款 : 粉丝 < 500                        -> 绿
  大爆款   : 24小时加购 > 800                  -> 红
  常规爆款 : 500 <= 粉丝 < 2000 且 加购 > 400   -> 黄
  未达标   : 其余                               -> 灰

额外加一列「可买性」（银河自己的预设标准，不是 SOP 原文）:
  银河三套预设（爆款/正常/低粉标准）全都带「到手价 >= 30」，低于 30 元的品
  在 SOP 调研表里算不过账（扣 5% 后单件利润只剩几毛），所以拿 30 元当门槛。

  到手价 >= 30           -> ✅ 可买
  到手价 < 30            -> ⚠️ 低价慎选
  price_suspect = 1      -> ❓ 价格存疑（接口自己标的价格不可信）
  deal_needs_coupon = 1  -> 文案追加「·需券」（要先去商品页领券才是这个价）

  另有一列「价格标记」记录原始标志（待核 / 需券）。
  注：/api/plaza 当前实际只返回 price_suspect；deal_needs_coupon 和
  original_price 前端有渲染但接口不返回（实测 900+ 条全为 null），
  脚本按「有就用」处理，字段上线后自动生效。

用法:
    set MC_TOKEN=ses.xxxx
    python fetch_plaza.py                          # SOP 全档：加购>400 的全部
    python fetch_plaza.py --max-followers 500 --min-cart 1 --limit 300
    python fetch_plaza.py --segment growth
"""
import argparse
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

BASE = "https://aiyinhe.art/api/plaza"
FILLS = {
    "低粉爆款": PatternFill("solid", fgColor="C6EFCE"),
    "大爆款": PatternFill("solid", fgColor="FFC7CE"),
    "常规爆款": PatternFill("solid", fgColor="FFEB9C"),
    "未达标": PatternFill("solid", fgColor="F2F2F2"),
}
SEGMENT_CN = {
    "premium": "高客单蓝海", "growth": "增长品(低粉爆单)", "mature": "成熟品",
    "popular": "热门品(腰部)", "potential": "潜力品", "watch": "观察",
}
OCEAN_CN = {"red": "红海", "blue": "普通蓝海", "": ""}


def read_flag(it, key):
    """读布尔型标记。接口会省略值为 null 的字段，所以拿不到就当 0。"""
    v = it.get(key)
    return 1 if v in (1, "1", True, "true") else 0


def classify(fans, cart):
    """SOP 4.0 分档，与 filter_scrape.py 保持一致。"""
    if fans is None or cart is None:
        return "未达标"
    if fans < 500:
        return "低粉爆款"
    if cart > 800:
        return "大爆款"
    if fans < 2000 and cart > 400:
        return "常规爆款"
    return "未达标"


# 可买性分档（银河预设标准要求 到手价 >= 30，见 README）
BUY_RANK = {"可买": 0, "待核": 1, "低价": 2}


def verdict(price, suspect, coupon):
    """返回 (可买性文案, 排序权重, 价格标记)。

    suspect = price_suspect，接口标了 1 表示这个价格不可信 -> 待核
    coupon  = deal_needs_coupon，接口标了 1 表示要先去商品页领券 -> 加「需券」后缀
    """
    if suspect or not price:
        text, rank = "❓ 价格存疑", BUY_RANK["待核"]
    elif price >= 30:
        text, rank = "✅ 可买", BUY_RANK["可买"]
    else:
        text, rank = "⚠️ 低价慎选", BUY_RANK["低价"]
    flags = ",".join(n for n, on in (("待核", bool(suspect)), ("需券", bool(coupon))) if on)
    return text + ("·需券" if coupon else ""), rank, flags


def fetch(token, params, limit=None, pause=0.35):
    """按页拉取，返回 dict 列表。遇到 429 按 retryAfterSeconds 等待。"""
    out, page = [], 1
    while True:
        q = dict(params, page=page, pageSize=50)
        url = BASE + "?" + urllib.parse.urlencode({k: v for k, v in q.items() if v not in ("", None)})
        req = urllib.request.Request(url, headers={
            "Authorization": f"Bearer {token}",
            "Accept": "application/json",
            "User-Agent": "Mozilla/5.0",
        })
        try:
            with urllib.request.urlopen(req, timeout=45) as r:
                body = json.loads(r.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code == 429:
                wait = 5
                try:
                    wait = int(json.loads(e.read().decode()).get("retryAfterSeconds", 5))
                except Exception:
                    pass
                print(f"  限流，等 {wait}s", file=sys.stderr)
                time.sleep(wait)
                continue
            raise
        if not body.get("success"):
            sys.exit(f"接口报错: {body.get('error')}")
        rows = body.get("data") or []
        out.extend(rows)
        total = (body.get("pagination") or {}).get("total", 0)
        print(f"  第 {page} 页 +{len(rows)} (累计 {len(out)}/{total})", file=sys.stderr)
        if not rows or len(out) >= total or (limit and len(out) >= limit):
            break
        page += 1
        time.sleep(pause)
    return out[:limit] if limit else out


def to_rows(items):
    data = []
    for it in items:
        fans = it.get("shop_followers")
        cart = it.get("add_cart_24h")
        sales = it.get("sales") or 0
        price = it.get("price")
        suspect = read_flag(it, "price_suspect")
        coupon = read_flag(it, "deal_needs_coupon")
        lvl = classify(fans, cart)
        ratio = round(cart / sales, 4) if sales else None
        buyable, _rank, flags = verdict(price, suspect, coupon)
        data.append([
            lvl, it.get("name") or "", price, cart,
            it.get("cart_delta_24h"), fans, sales, ratio,
            it.get("first_category_name") or "",
            SEGMENT_CN.get(it.get("strategy_segment") or "", it.get("strategy_segment") or ""),
            OCEAN_CN.get(it.get("ocean_type") or "", ""),
            flags, buyable, it.get("url") or "",
        ])
    assert len(data[0]) == len(TITLE), f"列数对不上: {len(data[0])} vs {len(TITLE)}"
    rank = {"低粉爆款": 0, "常规爆款": 1, "大爆款": 2, "未达标": 3}
    data.sort(key=lambda r: (rank[r[0]], -(r[3] or 0)))
    return data


TITLE = ["档位", "产品名", "到手价", "24小时加购", "日增加购", "店铺粉丝",
         "总销量", "加购/销量", "类目", "银河分层", "红蓝海", "价格标记", "可买性", "链接"]
WIDTHS = [10, 44, 9, 11, 10, 10, 10, 10, 10, 15, 9, 10, 13, 52]
URL_COL = TITLE.index("链接")
VERDICT_COL = TITLE.index("可买性")


def write_xlsx(data, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "选品广场"
    ws.append(TITLE)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for rec in data:
        ws.append(rec)
        for c in ws[ws.max_row]:
            c.fill = FILLS[rec[0]]
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for r in ws.iter_rows(min_row=2):
        cell = r[URL_COL]
        if cell.value:
            cell.hyperlink = str(cell.value)
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--token", default=os.environ.get("MC_TOKEN", ""),
                    help="aiyinhe.art 登录态 token（默认读环境变量 MC_TOKEN）")
    ap.add_argument("-o", "--output", default="选品广场-SOP分档.xlsx")
    ap.add_argument("--min-cart", type=int, default=401, help="24小时加购下限（SOP 常规档门槛 400）")
    ap.add_argument("--max-cart", type=int, default=0)
    ap.add_argument("--min-followers", type=int, default=0)
    ap.add_argument("--max-followers", type=int, default=0)
    ap.add_argument("--min-price", type=float, default=0)
    ap.add_argument("--max-price", type=float, default=0)
    ap.add_argument("--category", default="")
    ap.add_argument("--segment", default="")
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    if not a.token:
        sys.exit("缺少 token：python fetch_plaza.py --token ses.xxx  或先 set MC_TOKEN=ses.xxx")

    if os.path.exists(a.output):
        try:
            with open(a.output, "a+b"):
                pass
        except PermissionError:
            sys.exit(f"输出文件被占用，写不进去: {a.output}\n"
                     "大概率还开在 Excel 里，关掉再跑（或者用 -o 换个文件名）。")

    params = {
        "sort": "add_cart_24h", "order": "DESC",
        "min_cart": a.min_cart or "", "max_cart": a.max_cart or "",
        "min_followers": a.min_followers or "", "max_followers": a.max_followers or "",
        "min_price": a.min_price or "", "max_price": a.max_price or "",
        "category": a.category, "segment": a.segment,
    }
    print("拉取中…")
    items = fetch(a.token, params, limit=a.limit or None)
    if not items:
        sys.exit("没有拉到数据，检查筛选条件")

    data = to_rows(items)
    write_xlsx(data, a.output)

    from collections import Counter
    cnt = Counter(r[0] for r in data)
    buy = Counter(r[VERDICT_COL] for r in data)
    print(f"\n读入 {len(data)} 条 → {os.path.abspath(a.output)}")
    print("  低粉爆款 {低粉爆款} | 常规爆款 {常规爆款} | 大爆款 {大爆款} | 未达标 {未达标}".format(
        **{k: cnt.get(k, 0) for k in FILLS}))
    label = {c: sum(v for k, v in buy.items() if k.startswith(c)) for c in ("✅", "❓", "⚠️")}
    print(f"  可买 {label['✅']} | 价格存疑 {label['❓']} | 低价慎选(<30) {label['⚠️']}")

    if len(data) >= 50 and not any(read_flag(it, "deal_needs_coupon") for it in items):
        print("  ⚠️ deal_needs_coupon：本次 " + str(len(items)) + " 条里没有任何一条带这个字段。")
        print("     接口目前不返回它（original_price 同理），只有 price_suspect 是真的。")
        print("     脚本已按「有就用」接好，字段上线后自动生效，不用改代码。")

    top = [r for r in data if r[0] in ("低粉爆款", "常规爆款")][:15]
    if top:
        print("\n优先看（低粉爆款 + 常规爆款 TOP 15）:")
        for r in top:
            print(f"  [{r[0]}] {r[1][:30]:<32} 加购 {r[3]:>5} 粉丝 {r[5]:>6} ¥{r[2]}")


if __name__ == "__main__":
    main()
