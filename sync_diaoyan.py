# -*- coding: utf-8 -*-
"""把「二选表」里通过的行（F=是 且 G=是）推进「调研表」。

自动填 3 个值（从初选表搬）：
    竞品链接、竞品售价、24小时加购
不在自动范围内的（要么要你去查，要么要下单）：
    产品图/截图     <- 自己截图贴进去
    1688批发价/运费/包装 <- 1688 & 拼多多以图搜图后填
    同款竞品店铺数/高粉/低粉 <- 小红书以图搜图后数
    首篇爆文距今(天)  <- 小红书翻竞品最早的高赞笔记；或用 --first-seen-days 传
    货源1-5         <- 采购 3~5 家对比后填

用法:
    py sync_diaoyan.py                  # 同步
    py sync_diaoyan.py --dry-run
    py sync_diaoyan.py --first-seen-days 89
"""
import argparse
import os
import re
import sys

from openpyxl import load_workbook

CX = {"产品名": 4, "竞品链接": 5, "24小时加购": 6, "备注": 10}   # 初选表 1-based
EX = {"产品": 2, "拍摄": 6, "超越": 7}                          # 二选表
DY = {"竞品链接": 2, "竞品售价": 3, "24小时加购": 9, "首篇爆文距今": 16}  # 调研表


def cell(ws, row, col):
    v = ws.cell(row=row, column=col).value
    return "" if v is None else str(v).strip()


def price_from_note(note):
    """初选表备注形如 '[白牌] ¥59.9 销198000 增长品(低粉爆单)'，抠出第一个价格。

    注意别用 ^ 锚定 —— 备注前面现在带着 [白牌]/[品牌] 标记。
    """
    m = re.search(r"¥\s*([\d.]+)", str(note or ""))
    return float(m.group(1)) if m else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-f", "--file", default="选品SOP-4.0模板.xlsx")
    ap.add_argument("--first-seen-days", type=int, default=None,
                    help="竞品首篇爆文距今多少天（填这一列）")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    if not os.path.isfile(a.file):
        sys.exit(f"找不到文件: {a.file}")
    if not a.dry_run:
        try:
            with open(a.file, "a+b"):
                pass
        except PermissionError:
            sys.exit(f"文件被占用: {a.file}\n先保存再关闭 Excel，不然同步会丢你刚填的东西。")

    wb = load_workbook(a.file)
    cx, ex, dy = wb["初选表"], wb["二选表"], wb["调研表"]

    # 初选表：产品名 -> (链接, 加购, 售价)
    pool = {}
    for r in range(2, cx.max_row + 1):
        name = cell(cx, r, CX["产品名"])
        if name:
            pool[name] = (cell(cx, r, CX["竞品链接"]),
                          cx.cell(row=r, column=CX["24小时加购"]).value,
                          price_from_note(cell(cx, r, CX["备注"])))

    # 已经进过调研表的链接
    done = {cell(dy, r, DY["竞品链接"]) for r in range(2, dy.max_row + 1) if cell(dy, r, DY["竞品链接"])}
    next_row = max([r for r in range(2, dy.max_row + 1) if cell(dy, r, DY["竞品链接"])] + [1]) + 1

    todo = []
    for r in range(2, ex.max_row + 1):
        if cell(ex, r, EX["拍摄"]) == "是" and cell(ex, r, EX["超越"]) == "是":
            name = cell(ex, r, EX["产品"])
            link, cart, price = pool.get(name, ("", None, None))
            if link and link not in done:
                todo.append((r, name, link, cart, price))

    print(f"二选通过: {sum(1 for r in range(2, ex.max_row+1) if cell(ex,r,EX['拍摄'])=='是' and cell(ex,r,EX['超越'])=='是')} 行"
          f" | 调研表已有: {len(done)} 行 | 待推: {len(todo)} 行")

    if not todo:
        print("没有要推的")
        return

    for i, (r, name, link, cart, price) in enumerate(todo):
        row = next_row + i
        print(f"  调研表行{row:<3} <- 二选表行{r:<3}  {name[:28]}")
        print(f"        竞品链接 {link[:60]}")
        print(f"        竞品售价 {price} | 24小时加购 {cart}"
              + (f" | 首篇爆文距今 {a.first_seen_days}" if a.first_seen_days else ""))
        if not a.dry_run:
            c = dy.cell(row=row, column=DY["竞品链接"], value=link)
            c.hyperlink = link
            if price is not None:
                dy.cell(row=row, column=DY["竞品售价"], value=price)
            if cart is not None:
                dy.cell(row=row, column=DY["24小时加购"], value=cart)
            if a.first_seen_days is not None:
                dy.cell(row=row, column=DY["首篇爆文距今"], value=a.first_seen_days)

    if a.dry_run:
        print("\n--dry-run：什么都没写")
        return

    wb.save(a.file)
    print(f"\n已推进 {len(todo)} 行 → {os.path.abspath(a.file)}")
    print("  还差（要手工）：1688批发价 / 运费 / 包装 -> 单件利润自动算")
    print("                  同款竞品店铺数 / 高粉 / 低粉 -> 竞品总数&入场判断自动算")
    if not a.first_seen_days:
        print("  首篇爆文距今(天) 也还空着 -> 入场判断会显示「待查首篇爆文时间」")


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    main()
