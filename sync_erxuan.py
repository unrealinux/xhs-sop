# -*- coding: utf-8 -*-
"""把「初选表」里已经填好店铺名的行，同步到「二选表」。

SOP 流程（P19）：初选只填店铺名 -> 二选时「一个个复制初选里的店铺名，用手机 app
去查看他们的账号」。这个脚本干的就是那个「复制」，省得手抄。

只搬 2 列：店铺名、产品。其余留给公式和你的判断：
  笔记总数 / 链接评价数  <- 手机 app 里看竞品账号，手填
  账号权重(自动)          <- 公式，自己算
  拍摄能力能对齐吗 / 我能超越吗 <- 你填是/否
  结论(自动)              <- 公式，自己算

用法:
    py sync_erxuan.py                 # 同步
    py sync_erxuan.py --dry-run       # 只看会搬哪些，不写文件
    py sync_erxuan.py -f 我的表.xlsx
"""
import argparse
import os
import sys

from openpyxl import load_workbook

COL_CX = {"店铺名": 3, "产品名": 4}          # 初选表 1-based
COL_EX = {"店铺名": 1, "产品": 2}            # 二选表 1-based


def cell(ws, row, col):
    v = ws.cell(row=row, column=col).value
    return "" if v is None else str(v).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-f", "--file", default="选品SOP-4.0模板.xlsx")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    if not os.path.isfile(a.file):
        sys.exit(f"找不到文件: {a.file}")
    if not a.dry_run:
        try:
            with open(a.file, "a+b"):
                pass
        except PermissionError:
            sys.exit(f"文件被占用: {a.file}\n"
                     "Excel 还开着。**先保存再关闭**，不然同步会丢你刚填的东西。")

    wb = load_workbook(a.file)
    if "初选表" not in wb.sheetnames or "二选表" not in wb.sheetnames:
        sys.exit("这个文件里没有「初选表」或「二选表」")

    cx, ex = wb["初选表"], wb["二选表"]
    cap = ex.max_row - 1

    # 初选表里店铺名非空的行
    src = []
    for r in range(2, cx.max_row + 1):
        shop = cell(cx, r, COL_CX["店铺名"])
        if shop:
            src.append((r, shop, cell(cx, r, COL_CX["产品名"])))

    # 二选表里已有的（按 店铺名+产品 去重）
    have = set()
    next_row = 2
    for r in range(2, ex.max_row + 1):
        s, p = cell(ex, r, COL_EX["店铺名"]), cell(ex, r, COL_EX["产品"])
        if s or p:
            have.add((s, p))
            next_row = r + 1

    todo = [x for x in src if (x[1], x[2]) not in have]
    print(f"初选表里有店铺名: {len(src)} 行 | 二选表已有: {len(have)} 行 | 待搬: {len(todo)} 行")

    if not todo:
        print("没有要搬的（都已在二选表里，或者初选表还没填店铺名）")
        return

    if len(have) + len(todo) > cap:
        print(f"⚠️ 二选表只有 {cap} 行，装不下，多余的会被截掉")
        todo = todo[:cap - len(have)]

    for i, (r, shop, prod) in enumerate(todo):
        row = next_row + i
        print(f"  二选表行{row:<3} <- 初选表行{r:<3}  {shop}  |  {prod[:30]}")
        if not a.dry_run:
            ex.cell(row=row, column=COL_EX["店铺名"], value=shop)
            ex.cell(row=row, column=COL_EX["产品"], value=prod)

    if a.dry_run:
        print("\n--dry-run：什么都没写")
        return

    wb.save(a.file)
    print(f"\n已同步 {len(todo)} 行 → {os.path.abspath(a.file)}")
    print("  接下来：手机 app 打开这些店，把「笔记总数」「链接评价数」填进 C / D 列")


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    main()
